"""Excel比較結果を、1変更につき1行の読みやすいブックに保存する。"""

from datetime import date, datetime, time, timedelta
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils.exceptions import IllegalCharacterError


REPORTS_DIR = Path(__file__).resolve().parent / "reports"
HEADER_ROW = 8
DATA_START_ROW = HEADER_ROW + 1
HEADERS = ("No.", "種類", "シート名", "場所", "変更前", "変更後", "内容")
MAX_EXCEL_ROWS = 1_048_576
MAX_CELL_TEXT_LENGTH = 32_767

_ROW_COLORS = {
    "シート追加": "E2F0D9",
    "追加": "E2F0D9",
    "シート削除": "FCE4D6",
    "削除": "FCE4D6",
    "変更": "FFF2CC",
    "シート移動": "E4DFEC",
}
_NAVY = "203864"


class ReportWriteError(Exception):
    """Excelレポートを書き込めない場合の利用者向けエラー。"""


def _set_value(cell, value):
    """文字列は必ずテキストとして保存し、入力の数式を実行させない。"""
    if isinstance(value, str) and len(value) > MAX_CELL_TEXT_LENGTH:
        raise ReportWriteError(
            f"レポートの {cell.coordinate} に入る文字列がExcelの上限"
            f"（{MAX_CELL_TEXT_LENGTH:,}文字）を超えています。"
        )

    cell.value = value
    if isinstance(value, str):
        # =SUM(...) や #DIV/0! も数式・エラーではなく文字列として記録する。
        cell.data_type = "s"
    elif isinstance(value, datetime):
        cell.number_format = "yyyy/mm/dd hh:mm:ss"
    elif isinstance(value, date):
        cell.number_format = "yyyy/mm/dd"
    elif isinstance(value, time):
        cell.number_format = "hh:mm:ss"
    elif isinstance(value, timedelta):
        cell.number_format = "[h]:mm:ss"


def _report_rows(sheet_changes, cell_changes):
    for change in sheet_changes:
        kind = change["type"]
        old_position = change.get("old_position")
        new_position = change.get("new_position")
        description = {
            "シート追加": "シートを追加",
            "シート削除": "シートを削除",
            "シート移動": "共通シート間の順序変更",
        }[kind]
        yield (
            kind,
            change["sheet"],
            "シート全体",
            f"比較元の{old_position}番目" if old_position is not None else None,
            f"比較先の{new_position}番目" if new_position is not None else None,
            description,
        )

    for change in cell_changes:
        kind = change["type"]
        side = "旧" if kind == "削除" else "新"
        description = {
            "追加": "行を追加",
            "削除": "行を削除",
            "変更": "値を変更",
        }[kind]
        yield (
            kind,
            change["sheet"],
            f"{side} {change['cell']}",
            change["old"],
            change["new"],
            description,
        )


def _build_workbook(sheet_changes, cell_changes, old_file, new_file, created_at):
    count = len(sheet_changes) + len(cell_changes)
    if count > MAX_EXCEL_ROWS - HEADER_ROW:
        raise ReportWriteError(
            f"変更が多すぎるため、1枚のExcelシートに出力できません"
            f"（上限 {MAX_EXCEL_ROWS - HEADER_ROW:,}件）。"
        )

    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "変更点まとめ"
    sheet.sheet_view.showGridLines = False
    sheet.sheet_properties.pageSetUpPr.fitToPage = True
    sheet.sheet_properties.tabColor = _NAVY
    sheet.freeze_panes = "E9"

    sheet.merge_cells("A1:G1")
    _set_value(sheet["A1"], "Excel 変更点まとめ")
    sheet["A1"].font = Font(name="Yu Gothic", size=20, bold=True, color="FFFFFF")
    sheet["A1"].fill = PatternFill("solid", fgColor=_NAVY)
    sheet["A1"].alignment = Alignment(vertical="center")
    sheet.row_dimensions[1].height = 38

    metadata = (
        ("結果", f"{'変更あり' if count else '変更なし'}  |  "
         f"合計 {count:,}件（シート構成 {len(sheet_changes):,}件 / "
         f"セル {len(cell_changes):,}件）"),
        ("比較元", old_file.name),
        ("比較先", new_file.name),
        ("作成日時", created_at),
        ("見方", "緑＝追加 / 赤＝削除 / 黄＝変更 / 紫＝移動。旧＝比較元、新＝比較先の座標。"
         "値欄の空白＝空欄。行の追加・削除は非空セルごとに記録。"),
    )
    for row, (label, value) in enumerate(metadata, start=2):
        sheet.merge_cells(start_row=row, start_column=2, end_row=row, end_column=7)
        _set_value(sheet.cell(row, 1), label)
        _set_value(sheet.cell(row, 2), value)
        sheet.cell(row, 1).font = Font(name="Yu Gothic", bold=True, color=_NAVY)
        sheet.cell(row, 2).font = Font(name="Yu Gothic", size=11, color="243746")
        sheet.cell(row, 2).alignment = Alignment(horizontal="left", vertical="center", wrap_text=True)
        sheet.row_dimensions[row].height = 25 if row != 6 else 36
    sheet.row_dimensions[7].height = 10

    for column, header in enumerate(HEADERS, start=1):
        cell = sheet.cell(HEADER_ROW, column)
        _set_value(cell, header)
        cell.font = Font(name="Yu Gothic", bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor=_NAVY)
        cell.alignment = Alignment(vertical="center", wrap_text=True)
    sheet.row_dimensions[HEADER_ROW].height = 28

    border = Border(bottom=Side(style="hair", color="D5DBE5"))
    for index, values in enumerate(_report_rows(sheet_changes, cell_changes), start=1):
        row_number = HEADER_ROW + index
        fill = PatternFill("solid", fgColor=_ROW_COLORS[values[0]])
        for column, value in enumerate((index, *values), start=1):
            cell = sheet.cell(row_number, column)
            _set_value(cell, value)
            cell.font = Font(name="Yu Gothic", size=11, color="243746")
            cell.fill = fill
            cell.border = border
            cell.alignment = Alignment(vertical="top", wrap_text=True)

        # 長い値は折り返し表示。巨大な値でも行高が画面を占有しすぎないようにする。
        text_lines = max(
            (sum(max(1, (len(line) + 23) // 24) for line in str(value or "").split("\n"))
             for value in values),
            default=1,
        )
        sheet.row_dimensions[row_number].height = min(150, max(30, 16 * text_lines + 10))

    for column, width in zip("ABCDEFG", (12, 16, 24, 18, 38, 38, 34)):
        sheet.column_dimensions[column].width = width
    last_row = HEADER_ROW + count
    sheet.auto_filter.ref = f"A{HEADER_ROW}:G{last_row}"
    sheet.print_title_rows = f"1:{HEADER_ROW}"
    sheet.print_area = f"A1:G{last_row}"
    sheet.page_setup.orientation = "landscape"
    sheet.page_setup.paperSize = sheet.PAPERSIZE_A4
    sheet.page_setup.fitToWidth = 1
    sheet.page_setup.fitToHeight = 0
    return workbook


def write_report(sheet_changes, cell_changes, old_file: Path, new_file: Path,
                 output_path: Path | None = None) -> Path:
    """差分を新しい.xlsxへ保存し、絶対パスを返す。既存ファイルは上書きしない。"""
    created_at = datetime.now()
    output = Path(output_path) if output_path is not None else (
        REPORTS_DIR / f"変更点まとめ_{created_at:%Y%m%d_%H%M%S_%f}.xlsx"
    )
    workbook = None
    created_file = False
    try:
        output = output.expanduser().resolve()
        old_file, new_file = Path(old_file), Path(new_file)
        if output.suffix.lower() != ".xlsx":
            raise ReportWriteError("出力ファイルの拡張子は .xlsx を指定してください。")
        if output in (old_file.resolve(), new_file.resolve()):
            raise ReportWriteError("入力ファイルと同じ場所には出力できません。")
        if output.exists():
            raise ReportWriteError(f"出力先がすでに存在します: {output}")

        workbook = _build_workbook(
            sheet_changes, cell_changes, old_file, new_file, created_at
        )
        output.parent.mkdir(parents=True, exist_ok=True)
        # 排他作成により、確認後に同じ名前が作成された場合も上書きしない。
        with output.open("xb") as stream:
            created_file = True
            workbook.save(stream)
    except FileExistsError as exc:
        raise ReportWriteError(f"出力先がすでに存在します: {output}") from exc
    except (OSError, ValueError, TypeError, IllegalCharacterError, OverflowError) as exc:
        if created_file:
            try:
                output.unlink()
            except OSError:
                pass
        raise ReportWriteError(f"Excelレポートを保存できません: {output} ({exc})") from exc
    finally:
        if workbook is not None:
            workbook.close()
    return output

"""Excelレポートの内容、値の保持、出力先の保護を検証する。"""

from datetime import date, datetime
import os
from pathlib import Path
from tempfile import TemporaryDirectory
import unittest
from unittest.mock import patch

import openpyxl

import excel_report


class ExcelReportTests(unittest.TestCase):
    def setUp(self):
        self.temp = TemporaryDirectory(prefix="excel report ")
        self.addCleanup(self.temp.cleanup)
        self.directory = Path(self.temp.name)
        self.old = self.directory / "変更前.xlsx"
        self.new = self.directory / "変更後.xlsx"
        self.output = self.directory / "reports" / "変更点まとめ.xlsx"
        # レポート作成は入力ファイルを変更・再読込しない。
        self.old.write_bytes(b"original old file")
        self.new.write_bytes(b"original new file")

    def write_report(self, sheet_changes=None, cell_changes=None, output=None):
        return excel_report.write_report(
            sheet_changes or [], cell_changes or [], self.old, self.new,
            self.output if output is None else output,
        )

    def read_report(self, path=None, data_only=False):
        workbook = openpyxl.load_workbook(path or self.output, data_only=data_only)
        self.addCleanup(workbook.close)
        self.assertEqual(workbook.sheetnames, ["変更点まとめ"])
        worksheet = workbook.active
        headers = {
            cell.value: cell.column
            for cell in worksheet[8]
            if cell.value is not None
        }
        self.assertEqual(
            list(headers), ["No.", "種類", "シート名", "場所", "変更前", "変更後", "内容"]
        )
        records = [
            {name: worksheet.cell(row, column).value for name, column in headers.items()}
            for row in range(9, worksheet.max_row + 1)
            if worksheet.cell(row, headers["No."]).value is not None
        ]
        return worksheet, headers, records

    def test_every_change_has_one_row_in_input_order(self):
        sheet_changes = [
            {"type": "シート追加", "sheet": "新規", "new_position": 2},
            {"type": "シート削除", "sheet": "旧版", "old_position": 4},
            {"type": "シート移動", "sheet": "売上", "old_position": 3, "new_position": 1},
        ]
        cell_changes = [
            {"type": "追加", "sheet": "受注", "cell": "A3", "old": None, "new": "紅茶"},
            {"type": "削除", "sheet": "受注", "cell": "B4", "old": 100, "new": None},
            {"type": "変更", "sheet": "商品", "cell": "C5", "old": "未着手", "new": "完了"},
        ]

        result = self.write_report(sheet_changes, cell_changes)
        self.assertEqual(result, self.output.resolve())
        worksheet, _, records = self.read_report()
        self.assertEqual(worksheet.max_row, 14)
        self.assertEqual([row["No."] for row in records], list(range(1, 7)))
        self.assertEqual(
            [(row["種類"], row["シート名"]) for row in records],
            [(change["type"], change["sheet"]) for change in sheet_changes + cell_changes],
        )
        self.assertEqual(
            [(row["変更前"], row["変更後"]) for row in records[3:]],
            [(None, "紅茶"), (100, None), ("未着手", "完了")],
        )
        self.assertEqual([row["場所"] for row in records[3:]], ["新 A3", "旧 B4", "新 C5"])
        for row in records:
            self.assertIsInstance(row["内容"], str)
            self.assertTrue(row["内容"].strip())
        # シートの位置情報も行内で確認できること。
        for row, position in [(records[0], "2"), (records[1], "4")]:
            self.assertIn(position, " ".join(str(value) for value in row.values()))
        self.assertIn("3", str(records[2]["変更前"]))
        self.assertIn("1", str(records[2]["変更後"]))

    def test_metadata_and_navigation_make_saved_report_readable(self):
        self.write_report(cell_changes=[
            {"type": "変更", "sheet": "データ", "cell": "A1", "old": 10, "new": 20},
        ])
        worksheet, _, _ = self.read_report()
        metadata = " ".join(
            str(cell.value)
            for row in worksheet.iter_rows(min_row=1, max_row=6)
            for cell in row
            if cell.value is not None
        )
        self.assertIn(self.old.name, metadata)
        self.assertIn(self.new.name, metadata)
        self.assertIn("変更あり", metadata)
        self.assertIn("空欄", metadata)
        self.assertEqual(worksheet.auto_filter.ref, "A8:G9")
        self.assertEqual(worksheet.freeze_panes, "E9")
        self.assertTrue(worksheet.cell(9, 7).alignment.wrap_text)

    def test_no_changes_still_creates_report_without_fake_change(self):
        self.write_report()
        worksheet, _, records = self.read_report()
        self.assertEqual(records, [])
        metadata = " ".join(
            str(cell.value)
            for row in worksheet.iter_rows(min_row=1, max_row=6)
            for cell in row
            if cell.value is not None
        )
        self.assertIn("変更なし", metadata)

    def test_formulas_and_excel_error_strings_remain_literal_text(self):
        changes = [{
            "type": "変更", "sheet": "=1+1", "cell": "A1",
            "old": "=SUM(A1:A2)", "new": "#REF!",
        }]
        self.write_report(cell_changes=changes)
        for data_only in (False, True):
            with self.subTest(data_only=data_only):
                worksheet, headers, records = self.read_report(data_only=data_only)
                self.assertEqual(records[0]["変更前"], "=SUM(A1:A2)")
                self.assertEqual(records[0]["変更後"], "#REF!")
                self.assertEqual(records[0]["シート名"], "=1+1")
                for name in ("変更前", "変更後", "シート名"):
                    self.assertEqual(worksheet.cell(9, headers[name]).data_type, "s")
                self.assertFalse(any(
                    cell.data_type == "f"
                    for row in worksheet
                    for cell in row
                ))

    def test_blank_zero_boolean_and_dates_keep_meaningful_values(self):
        timestamp = datetime(2026, 9, 23, 12, 30)
        changes = [
            {"type": "変更", "sheet": "データ", "cell": "A1", "old": None, "new": 0},
            {"type": "変更", "sheet": "データ", "cell": "A2", "old": False, "new": True},
            {"type": "変更", "sheet": "データ", "cell": "A3",
             "old": date(2026, 9, 22), "new": timestamp},
        ]
        self.write_report(cell_changes=changes)
        worksheet, headers, records = self.read_report()
        self.assertIsNone(records[0]["変更前"])
        self.assertEqual(records[0]["変更後"], 0)
        self.assertIs(type(records[0]["変更後"]), int)
        self.assertIs(records[1]["変更前"], False)
        self.assertIs(records[1]["変更後"], True)
        self.assertEqual(records[2]["変更前"].date(), date(2026, 9, 22))
        self.assertEqual(records[2]["変更後"], timestamp)
        self.assertTrue(worksheet.cell(11, headers["変更前"]).is_date)
        self.assertTrue(worksheet.cell(11, headers["変更後"]).is_date)

    def test_same_position_sheet_move_is_explained(self):
        self.write_report(sheet_changes=[
            {"type": "シート移動", "sheet": "売上", "old_position": 2, "new_position": 2},
        ])
        _, _, records = self.read_report()
        self.assertEqual(len(records), 1)
        self.assertIn("順序", records[0]["内容"])

    def test_default_output_uses_descriptive_name_in_reports_directory(self):
        report_directory = self.directory / "default reports"
        with patch.object(excel_report, "REPORTS_DIR", report_directory):
            output = excel_report.write_report([], [], self.old, self.new)
        self.assertEqual(output.parent, report_directory.resolve())
        self.assertRegex(output.name, r"^変更点まとめ_\d{8}_\d{6}_\d{6}\.xlsx$")
        self.assertTrue(output.is_file())

    def test_relative_output_is_resolved_from_current_directory(self):
        original_directory = Path.cwd()
        try:
            os.chdir(self.directory)
            output = self.write_report(output=Path("nested folder") / "report.xlsx")
        finally:
            os.chdir(original_directory)
        self.assertEqual(output, (self.directory / "nested folder" / "report.xlsx").resolve())
        self.assertTrue(output.is_file())

    def test_existing_output_is_never_overwritten(self):
        self.output.parent.mkdir()
        original = b"do not overwrite this report"
        self.output.write_bytes(original)
        with self.assertRaises(excel_report.ReportWriteError):
            self.write_report()
        self.assertEqual(self.output.read_bytes(), original)
        self.assertEqual(list(self.output.parent.iterdir()), [self.output])

    def test_input_paths_are_never_overwritten(self):
        for input_path in (self.old, self.new):
            with self.subTest(input_path=input_path):
                original = input_path.read_bytes()
                with self.assertRaises(excel_report.ReportWriteError):
                    self.write_report(output=input_path)
                self.assertEqual(input_path.read_bytes(), original)

    def test_non_xlsx_output_is_rejected(self):
        output = self.directory / "report.csv"
        with self.assertRaises(excel_report.ReportWriteError):
            self.write_report(output=output)
        self.assertFalse(output.exists())

    def test_file_in_place_of_parent_directory_is_friendly_error(self):
        parent = self.directory / "not a directory"
        parent.write_bytes(b"keep this file")
        with self.assertRaises(excel_report.ReportWriteError):
            self.write_report(output=parent / "report.xlsx")
        self.assertEqual(parent.read_bytes(), b"keep this file")

    def test_failed_save_removes_partial_output(self):
        def fail_save(workbook, target):
            if hasattr(target, "write"):
                target.write(b"partial workbook")
            else:
                Path(target).write_bytes(b"partial workbook")
            raise OSError("テスト用の書き込み失敗")

        with patch.object(openpyxl.Workbook, "save", autospec=True, side_effect=fail_save):
            with self.assertRaises(excel_report.ReportWriteError):
                self.write_report()
        self.assertFalse(self.output.exists())
        if self.output.parent.exists():
            self.assertEqual(list(self.output.parent.iterdir()), [])
        self.assertEqual(self.old.read_bytes(), b"original old file")
        self.assertEqual(self.new.read_bytes(), b"original new file")


if __name__ == "__main__":
    unittest.main()

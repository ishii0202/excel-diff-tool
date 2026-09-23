"""比較ロジックとコマンドラインの回帰テスト。"""

from contextlib import redirect_stderr, redirect_stdout
from io import StringIO
import os
from pathlib import Path
import subprocess
import sys
from tempfile import TemporaryDirectory
import unittest
from unittest.mock import patch
from zipfile import ZipFile

import openpyxl

import excel_diff


ROOT = Path(__file__).resolve().parents[1]
SCRIPT = ROOT / "excel_diff.py"


def make_workbook(rows, sheets=None):
    wb = openpyxl.Workbook()
    wb.active.title = "データ"
    for row in rows:
        wb.active.append(row)
    for sheet in sheets or []:
        wb.create_sheet(sheet)
    return wb


class ComparisonTests(unittest.TestCase):
    def compare_rows(self, old_rows, new_rows):
        old = make_workbook(old_rows)
        new = make_workbook(new_rows)
        self.addCleanup(old.close)
        self.addCleanup(new.close)
        return excel_diff.compare_sheet("データ", old.active, new.active)

    def test_identical_values_and_formulas_have_no_changes(self):
        rows = [("商品", "数量", "金額"), ("紅茶", 2, "=B2*500"), (None, 0, False)]
        self.assertEqual(self.compare_rows(rows, rows), [])

    def test_inserted_row_does_not_change_following_rows(self):
        changes = self.compare_rows(
            [("A", 10), ("C", 30)],
            [("A", 10), ("B", 20), ("C", 30)],
        )
        self.assertEqual(
            [(c["type"], c["cell"], c["new"]) for c in changes],
            [("追加", "A2", "B"), ("追加", "B2", 20)],
        )

    def test_deleted_row_does_not_change_following_rows(self):
        changes = self.compare_rows(
            [("A", 10), ("B", 20), ("C", 30)],
            [("A", 10), ("C", 30)],
        )
        self.assertEqual(
            [(c["type"], c["cell"], c["old"]) for c in changes],
            [("削除", "A2", "B"), ("削除", "B2", 20)],
        )

    def test_values_and_formula_text_changes_are_reported(self):
        changes = self.compare_rows(
            [("数量", "合計", "メモ"), (2, "=A2*10", None)],
            [("数量", "合計", "メモ"), (3, "=A2*20", "更新")],
        )
        self.assertEqual(
            [(c["type"], c["cell"], c["old"], c["new"]) for c in changes],
            [("変更", "A2", 2, 3), ("変更", "B2", "=A2*10", "=A2*20"),
             ("変更", "C2", None, "更新")],
        )

    def test_extra_replacement_rows_are_added_and_deleted(self):
        old_rows = [("見出し",), ("旧商品",), ("終わり",)]
        new_rows = [("見出し",), ("新商品",), ("追加商品",), ("終わり",)]
        added = self.compare_rows(old_rows, new_rows)
        removed = self.compare_rows(new_rows, old_rows)
        self.assertEqual([(c["type"], c["cell"]) for c in added],
                         [("変更", "A2"), ("追加", "A3")])
        self.assertEqual([(c["type"], c["cell"]) for c in removed],
                         [("変更", "A2"), ("削除", "A3")])

    def sheet_changes(self, old_names, new_names):
        books = []
        for names in (old_names, new_names):
            wb = openpyxl.Workbook()
            wb.active.title = names[0]
            for name in names[1:]:
                wb.create_sheet(name)
            books.append(wb)
            self.addCleanup(wb.close)
        return excel_diff.detect_sheet_changes(*books)

    def test_sheet_addition_does_not_move_following_sheets(self):
        self.assertEqual(self.sheet_changes(["A", "B", "C"], ["A", "新規", "B", "C"]),
                         [{"type": "シート追加", "sheet": "新規", "new_position": 2}])

    def test_sheet_deletion_does_not_move_following_sheets(self):
        self.assertEqual(self.sheet_changes(["A", "旧", "B", "C"], ["A", "B", "C"]),
                         [{"type": "シート削除", "sheet": "旧", "old_position": 2}])

    def test_reordering_keeps_relative_order_change_at_same_position(self):
        changes = self.sheet_changes(["A", "B", "C"], ["C", "B", "A"])
        self.assertIn(
            {"type": "シート移動", "sheet": "B", "old_position": 2, "new_position": 2},
            changes,
        )

    def test_same_position_order_change_has_explanatory_output(self):
        output = StringIO()
        with redirect_stdout(output):
            excel_diff.print_results([
                {"type": "シート移動", "sheet": "B", "old_position": 3, "new_position": 3}
            ], [])
        self.assertIn("シート名: B", output.getvalue())
        self.assertIn("場所: 3番目（共通シート間の順序変更）", output.getvalue())
        self.assertNotIn("変更前:", output.getvalue())
        self.assertNotIn("変更後:", output.getvalue())

    def test_additions_and_deletion_do_not_hide_relative_reordering(self):
        changes = self.sheet_changes(["X", "A", "B"], ["Y", "Z", "B", "A"])
        self.assertIn(
            {"type": "シート移動", "sheet": "B", "old_position": 3, "new_position": 3},
            changes,
        )
        self.assertEqual(
            {change["sheet"] for change in changes if change["type"] == "シート追加"},
            {"Y", "Z"},
        )
        self.assertEqual(
            {change["sheet"] for change in changes if change["type"] == "シート削除"},
            {"X"},
        )

    def test_sheet_addition_deletion_and_reordering(self):
        changes = self.sheet_changes(["A", "B", "C", "削除"], ["C", "A", "B", "追加"])
        self.assertEqual(changes, [
            {"type": "シート追加", "sheet": "追加", "new_position": 4},
            {"type": "シート削除", "sheet": "削除", "old_position": 4},
            {"type": "シート移動", "sheet": "C", "old_position": 3, "new_position": 1},
        ])


class CommandLineTests(unittest.TestCase):
    def setUp(self):
        self.temp = TemporaryDirectory(prefix="excel diff ")
        self.addCleanup(self.temp.cleanup)
        self.directory = Path(self.temp.name)
        self.old = self.directory / "old book.xlsx"
        self.new = self.directory / "new book.xlsx"
        self.report = self.directory / "変更点まとめ.xlsx"
        for path in (self.old, self.new):
            wb = make_workbook([("商品", "数量"), ("紅茶", 2)])
            wb.save(path)
            wb.close()

    def run_cli(self, *args, output=None):
        return subprocess.run(
            [sys.executable, str(SCRIPT), "--output",
             str(output if output is not None else self.report), *map(str, args)],
            cwd=self.directory,
            env={**os.environ, "PYTHONIOENCODING": "utf-8"},
            capture_output=True, text=True, encoding="utf-8", check=False,
        )

    def test_help(self):
        result = self.run_cli("--help")
        self.assertEqual(result.returncode, 0)
        self.assertIn("old_file", result.stdout)
        self.assertIn("--output", result.stdout)

    def test_paths_with_spaces_and_different_working_directory(self):
        result = self.run_cli(self.old, self.new)
        self.assertEqual(result.returncode, 0, result.stderr)
        self.assertEqual(result.stdout.splitlines()[0], "変更なし")
        self.assertIn(str(self.report), result.stdout)
        self.assertTrue(self.report.is_file())

    def test_default_examples_from_different_working_directory(self):
        result = self.run_cli()
        self.assertEqual(result.returncode, 0, result.stderr)
        self.assertIn("変更あり", result.stdout)

    def test_custom_output_directory_and_input_files_stay_unchanged(self):
        old_bytes = self.old.read_bytes()
        new_bytes = self.new.read_bytes()
        output = self.directory / "review files" / "比較結果.xlsx"
        result = self.run_cli(self.old, self.new, output=output)
        self.assertEqual(result.returncode, 0, result.stderr)
        self.assertTrue(output.is_file())
        self.assertIn(str(output), result.stdout)
        self.assertEqual(self.old.read_bytes(), old_bytes)
        self.assertEqual(self.new.read_bytes(), new_bytes)

    def test_existing_report_is_not_overwritten(self):
        self.report.write_bytes(b"keep this report")
        result = self.run_cli(self.old, self.new)
        self.assertEqual(result.returncode, 1)
        self.assertEqual(self.report.read_bytes(), b"keep this report")
        self.assertNotIn("レポート保存先:", result.stdout)
        self.assertNotIn("Traceback", result.stderr)

    def test_output_cannot_replace_input(self):
        original = self.old.read_bytes()
        result = self.run_cli(self.old, self.new, output=self.old)
        self.assertEqual(result.returncode, 1)
        self.assertEqual(self.old.read_bytes(), original)
        self.assertNotIn("Traceback", result.stderr)

    def test_invalid_report_extension_is_friendly_error(self):
        result = self.run_cli(self.old, self.new, output=self.directory / "report.csv")
        self.assertEqual(result.returncode, 1)
        self.assertNotIn("Traceback", result.stderr)
        self.assertFalse((self.directory / "report.csv").exists())

    def test_one_argument_is_usage_error(self):
        result = self.run_cli(self.old)
        self.assertEqual(result.returncode, 2)
        self.assertIn("両方指定", result.stderr)

    def test_missing_file_is_friendly_error(self):
        result = self.run_cli(self.old, self.directory / "missing.xlsx")
        self.assertEqual(result.returncode, 1)
        self.assertIn("ファイルが見つかりません", result.stderr)
        self.assertIn("missing.xlsx", result.stderr)
        self.assertNotIn("Traceback", result.stderr)

    def test_invalid_excel_file_is_friendly_error(self):
        self.new.write_text("This is not an Excel workbook.")
        result = self.run_cli(self.old, self.new)
        self.assertEqual(result.returncode, 1)
        self.assertIn("Excelファイルを読み込めません", result.stderr)
        self.assertNotIn("Traceback", result.stderr)

    def test_zip_without_workbook_is_friendly_error(self):
        with ZipFile(self.new, "w") as archive:
            archive.writestr("unrelated.txt", "No workbook here")
        result = self.run_cli(self.old, self.new)
        self.assertEqual(result.returncode, 1)
        self.assertIn("Excelファイルを読み込めません", result.stderr)
        self.assertNotIn("Traceback", result.stderr)

    def test_import_has_no_output_or_file_access(self):
        code = (
            "import sys; from unittest.mock import patch; "
            f"sys.path.insert(0, {str(ROOT)!r}); "
            "patcher = patch('openpyxl.load_workbook', side_effect=AssertionError('unexpected load')); "
            "patcher.start(); import excel_diff"
        )
        result = subprocess.run([sys.executable, "-c", code], cwd=self.directory,
                                capture_output=True, text=True, check=False)
        self.assertEqual(result.returncode, 0, result.stderr)
        self.assertEqual(result.stdout, "")
        self.assertEqual(result.stderr, "")

    def test_first_workbook_closes_when_second_load_fails(self):
        old = make_workbook([])
        with patch.object(old, "close") as close, patch.object(
            excel_diff, "load_workbook", side_effect=[old, excel_diff.WorkbookReadError("壊れたブック")]
        ), redirect_stderr(StringIO()):
            self.assertEqual(excel_diff.main([str(self.old), str(self.new)]), 1)
        close.assert_called_once_with()

    def test_both_workbooks_close_when_comparison_fails(self):
        old = make_workbook([])
        new = make_workbook([])
        with patch.object(old, "close") as close_old, patch.object(new, "close") as close_new, patch.object(
            excel_diff, "load_workbook", side_effect=[old, new]
        ), patch.object(excel_diff, "compare_sheet", side_effect=RuntimeError("comparison failed")):
            with self.assertRaisesRegex(RuntimeError, "comparison failed"):
                excel_diff.main([str(self.old), str(self.new)])
        close_old.assert_called_once_with()
        close_new.assert_called_once_with()


if __name__ == "__main__":
    unittest.main()

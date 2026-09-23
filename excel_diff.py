"""Excelブックを比較し、結果を表示・レポートに保存する。"""

import argparse
from contextlib import ExitStack, closing
from difflib import SequenceMatcher
from datetime import date, datetime, time, timedelta
from pathlib import Path
import sys
from zipfile import BadZipFile

import openpyxl
from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils.exceptions import IllegalCharacterError, InvalidFileException


EXAMPLES_DIR = Path(__file__).resolve().parent / "examples"


class WorkbookReadError(Exception):
    """入力ブックを読み込めない場合の利用者向けエラー。"""


def load_workbook(path):
    """数式を文字列のまま読み込み、入力エラーにファイル名を添える。"""
    if not path.is_file():
        raise WorkbookReadError(f"ファイルが見つかりません: {path}")

    try:
        return openpyxl.load_workbook(path, data_only=False)
    except (
        OSError, ValueError, KeyError, IndexError, TypeError,
        SyntaxError, EOFError, BadZipFile, InvalidFileException,
    ) as exc:
        raise WorkbookReadError(
            f"Excelファイルを読み込めません: {path} ({exc})"
        ) from exc


# ============================================================
# 共通処理
# ============================================================

def display_value(value):
    """Noneを見やすく表示する"""
    if value is None:
        return "空欄"

    return value


def get_rows(ws, max_col):
    """
    ワークシートを行単位で取得する。

    例:
    [
        ("項目1", "説明1", 100),
        ("項目2", "説明2", 200),
        ...
    ]
    """
    rows = []

    for row in range(1, ws.max_row + 1):
        row_values = tuple(
            ws.cell(row, col).value
            for col in range(1, max_col + 1)
        )

        rows.append(row_values)

    return rows


# ============================================================
# シートの追加・削除・移動
# ============================================================

def detect_sheet_changes(old_wb, new_wb):
    old_names = old_wb.sheetnames
    new_names = new_wb.sheetnames

    old_set = set(old_names)
    new_set = set(new_names)

    changes = []

    # --------------------------------------------------------
    # シート追加
    # --------------------------------------------------------

    for index, sheet_name in enumerate(new_names):

        if sheet_name not in old_set:
            changes.append({
                "type": "シート追加",
                "sheet": sheet_name,
                "new_position": index + 1
            })

    # --------------------------------------------------------
    # シート削除
    # --------------------------------------------------------

    for index, sheet_name in enumerate(old_names):

        if sheet_name not in new_set:
            changes.append({
                "type": "シート削除",
                "sheet": sheet_name,
                "old_position": index + 1
            })

    # --------------------------------------------------------
    # 両方に存在するシートだけ抽出
    #
    # 例:
    #
    # 旧
    # A B C
    #
    # 新
    # A X B C
    #
    # ↓
    #
    # common_old = A B C
    # common_new = A B C
    #
    # なのでB,Cは移動扱いにならない
    # --------------------------------------------------------

    common_old = [
        name
        for name in old_names
        if name in new_set
    ]

    common_new = [
        name
        for name in new_names
        if name in old_set
    ]

    # --------------------------------------------------------
    # SequenceMatcherの一致ブロックを利用して
    # 「順序を維持しているシート」を探す
    # --------------------------------------------------------

    matcher = SequenceMatcher(
        None,
        common_old,
        common_new,
        autojunk=False
    )

    stable_sheets = set()

    for block in matcher.get_matching_blocks():

        for i in range(block.size):
            sheet_name = common_old[block.a + i]

            stable_sheets.add(sheet_name)

    # --------------------------------------------------------
    # 共通シートなのに一致ブロックから外れたもの
    # = 実際に移動した可能性が高いシート
    # --------------------------------------------------------

    common_set = old_set & new_set

    for sheet_name in old_names:

        if sheet_name not in common_set:
            continue

        if sheet_name in stable_sheets:
            continue

        old_position = old_names.index(sheet_name) + 1
        new_position = new_names.index(sheet_name) + 1

        changes.append({
            "type": "シート移動",
            "sheet": sheet_name,
            "old_position": old_position,
            "new_position": new_position
        })

    return changes


# ============================================================
# シート内部の比較
# ============================================================

def compare_sheet(sheet_name, old_ws, new_ws):
    changes = []

    max_col = max(
        old_ws.max_column,
        new_ws.max_column
    )

    old_rows = get_rows(old_ws, max_col)
    new_rows = get_rows(new_ws, max_col)

    # --------------------------------------------------------
    # 行単位で比較
    #
    # これによって途中に行が追加されても
    # 後続行を全部変更扱いしない
    # --------------------------------------------------------

    matcher = SequenceMatcher(
        None,
        old_rows,
        new_rows,
        autojunk=False
    )

    for tag, i1, i2, j1, j2 in matcher.get_opcodes():

        # ----------------------------------------------------
        # 同一部分
        # ----------------------------------------------------

        if tag == "equal":
            continue

        # ----------------------------------------------------
        # 行追加
        # ----------------------------------------------------

        if tag == "insert":

            for new_index in range(j1, j2):

                new_row_number = new_index + 1

                for col in range(1, max_col + 1):

                    new_value = new_rows[new_index][col - 1]

                    # 空欄は出力しない
                    if new_value is None:
                        continue

                    cell = new_ws.cell(
                        new_row_number,
                        col
                    ).coordinate

                    changes.append({
                        "type": "追加",
                        "sheet": sheet_name,
                        "cell": cell,
                        "old": None,
                        "new": new_value
                    })

        # ----------------------------------------------------
        # 行削除
        # ----------------------------------------------------

        elif tag == "delete":

            for old_index in range(i1, i2):

                old_row_number = old_index + 1

                for col in range(1, max_col + 1):

                    old_value = old_rows[old_index][col - 1]

                    if old_value is None:
                        continue

                    cell = old_ws.cell(
                        old_row_number,
                        col
                    ).coordinate

                    changes.append({
                        "type": "削除",
                        "sheet": sheet_name,
                        "cell": cell,
                        "old": old_value,
                        "new": None
                    })

        # ----------------------------------------------------
        # 行の中身が変更された
        # ----------------------------------------------------

        elif tag == "replace":

            old_count = i2 - i1
            new_count = j2 - j1

            pair_count = min(
                old_count,
                new_count
            )

            # ------------------------------------------------
            # 対応する行同士をセル単位で比較
            # ------------------------------------------------

            for offset in range(pair_count):

                old_index = i1 + offset
                new_index = j1 + offset

                for col in range(1, max_col + 1):

                    old_value = old_rows[old_index][col - 1]
                    new_value = new_rows[new_index][col - 1]

                    if old_value == new_value:
                        continue

                    cell = new_ws.cell(
                        new_index + 1,
                        col
                    ).coordinate

                    changes.append({
                        "type": "変更",
                        "sheet": sheet_name,
                        "cell": cell,
                        "old": old_value,
                        "new": new_value
                    })

            # ------------------------------------------------
            # 旧側の方が行数が多い
            # → 余った行は削除
            # ------------------------------------------------

            if old_count > new_count:

                for old_index in range(
                    i1 + pair_count,
                    i2
                ):

                    for col in range(1, max_col + 1):

                        old_value = old_rows[old_index][col - 1]

                        if old_value is None:
                            continue

                        cell = old_ws.cell(
                            old_index + 1,
                            col
                        ).coordinate

                        changes.append({
                            "type": "削除",
                            "sheet": sheet_name,
                            "cell": cell,
                            "old": old_value,
                            "new": None
                        })

            # ------------------------------------------------
            # 新側の方が行数が多い
            # → 余った行は追加
            # ------------------------------------------------

            elif new_count > old_count:

                for new_index in range(
                    j1 + pair_count,
                    j2
                ):

                    for col in range(1, max_col + 1):

                        new_value = new_rows[new_index][col - 1]

                        if new_value is None:
                            continue

                        cell = new_ws.cell(
                            new_index + 1,
                            col
                        ).coordinate

                        changes.append({
                            "type": "追加",
                            "sheet": sheet_name,
                            "cell": cell,
                            "old": None,
                            "new": new_value
                        })

    return changes


# ============================================================
# 結果表示
# ============================================================

def print_results(sheet_changes, cell_changes):

    if not sheet_changes and not cell_changes:
        print("変更なし")
        return

    print("変更あり")
    print()

    # --------------------------------------------------------
    # シート変更
    # --------------------------------------------------------

    for change in sheet_changes:

        if change["type"] == "シート追加":

            print("種類: シート追加")
            print(f"シート名: {change['sheet']}")
            print(f"場所: {change['new_position']}番目")
            print()

        elif change["type"] == "シート削除":

            print("種類: シート削除")
            print(f"シート名: {change['sheet']}")
            print(f"変更前: {change['old_position']}番目")
            print()

        elif change["type"] == "シート移動":

            print("種類: シート移動")
            print(f"シート名: {change['sheet']}")
            if change["old_position"] == change["new_position"]:
                print(
                    f"場所: {change['new_position']}番目"
                    "（共通シート間の順序変更）"
                )
            else:
                print(f"変更前: {change['old_position']}番目")
                print(f"変更後: {change['new_position']}番目")
            print()

    # --------------------------------------------------------
    # セル変更
    # --------------------------------------------------------

    for change in cell_changes:

        print(f"種類: {change['type']}")
        print(f"シート名: {change['sheet']}")
        print(f"場所: {change['cell']}")
        print(
            f"変更前: "
            f"{display_value(change['old'])}"
        )
        print(
            f"変更後: "
            f"{display_value(change['new'])}"
        )
        print()


# ============================================================
# Excelレポートの作成
# ============================================================

REPORTS_DIR = Path(__file__).resolve().parent / "reports"
HEADER_ROW = 8
DATA_START_ROW = HEADER_ROW + 1
HEADERS = ("No.", "種類", "シート名", "場所", "変更前", "変更後", "内容")
MAX_EXCEL_ROWS = 1_048_576
MAX_CELL_TEXT_LENGTH = 32_767

_ROW_COLORS = {
    "シート追加": "E2F0D9",
    "追加": "E2F0D9",
    "シート削除": "FCE4D6",
    "削除": "FCE4D6",
    "変更": "FFF2CC",
    "シート移動": "E4DFEC",
}
_NAVY = "203864"


class ReportWriteError(Exception):
    """Excelレポートを書き込めない場合の利用者向けエラー。"""


def _set_value(cell, value):
    """文字列は必ずテキストとして保存し、入力の数式を実行させない。"""
    if isinstance(value, str) and len(value) > MAX_CELL_TEXT_LENGTH:
        raise ReportWriteError(
            f"レポートの {cell.coordinate} に入る文字列がExcelの上限"
            f"（{MAX_CELL_TEXT_LENGTH:,}文字）を超えています。"
        )

    cell.value = value
    if isinstance(value, str):
        # =SUM(...) や #DIV/0! も数式・エラーではなく文字列として記録する。
        cell.data_type = "s"
    elif isinstance(value, datetime):
        cell.number_format = "yyyy/mm/dd hh:mm:ss"
    elif isinstance(value, date):
        cell.number_format = "yyyy/mm/dd"
    elif isinstance(value, time):
        cell.number_format = "hh:mm:ss"
    elif isinstance(value, timedelta):
        cell.number_format = "[h]:mm:ss"


def _report_rows(sheet_changes, cell_changes):
    for change in sheet_changes:
        kind = change["type"]
        old_position = change.get("old_position")
        new_position = change.get("new_position")
        description = {
            "シート追加": "シートを追加",
            "シート削除": "シートを削除",
            "シート移動": "共通シート間の順序変更",
        }[kind]
        yield (
            kind,
            change["sheet"],
            "シート全体",
            f"比較元の{old_position}番目" if old_position is not None else None,
            f"比較先の{new_position}番目" if new_position is not None else None,
            description,
        )

    for change in cell_changes:
        kind = change["type"]
        side = "旧" if kind == "削除" else "新"
        description = {
            "追加": "行を追加",
            "削除": "行を削除",
            "変更": "値を変更",
        }[kind]
        yield (
            kind,
            change["sheet"],
            f"{side} {change['cell']}",
            change["old"],
            change["new"],
            description,
        )


def _build_workbook(sheet_changes, cell_changes, old_file, new_file, created_at):
    count = len(sheet_changes) + len(cell_changes)
    if count > MAX_EXCEL_ROWS - HEADER_ROW:
        raise ReportWriteError(
            f"変更が多すぎるため、1枚のExcelシートに出力できません"
            f"（上限 {MAX_EXCEL_ROWS - HEADER_ROW:,}件）。"
        )

    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "変更点まとめ"
    sheet.sheet_view.showGridLines = False
    sheet.sheet_properties.pageSetUpPr.fitToPage = True
    sheet.sheet_properties.tabColor = _NAVY
    sheet.freeze_panes = "E9"

    sheet.merge_cells("A1:G1")
    _set_value(sheet["A1"], "Excel 変更点まとめ")
    sheet["A1"].font = Font(name="Yu Gothic", size=20, bold=True, color="FFFFFF")
    sheet["A1"].fill = PatternFill("solid", fgColor=_NAVY)
    sheet["A1"].alignment = Alignment(vertical="center")
    sheet.row_dimensions[1].height = 38

    metadata = (
        ("結果", f"{'変更あり' if count else '変更なし'}  |  "
         f"合計 {count:,}件（シート構成 {len(sheet_changes):,}件 / "
         f"セル {len(cell_changes):,}件）"),
        ("比較元", old_file.name),
        ("比較先", new_file.name),
        ("作成日時", created_at),
        ("見方", "緑＝追加 / 赤＝削除 / 黄＝変更 / 紫＝移動。旧＝比較元、新＝比較先の座標。"
         "値欄の空白＝空欄。行の追加・削除は非空セルごとに記録。"),
    )
    for row, (label, value) in enumerate(metadata, start=2):
        sheet.merge_cells(start_row=row, start_column=2, end_row=row, end_column=7)
        _set_value(sheet.cell(row, 1), label)
        _set_value(sheet.cell(row, 2), value)
        sheet.cell(row, 1).font = Font(name="Yu Gothic", bold=True, color=_NAVY)
        sheet.cell(row, 2).font = Font(name="Yu Gothic", size=11, color="243746")
        sheet.cell(row, 2).alignment = Alignment(horizontal="left", vertical="center", wrap_text=True)
        sheet.row_dimensions[row].height = 25 if row != 6 else 36
    sheet.row_dimensions[7].height = 10

    for column, header in enumerate(HEADERS, start=1):
        cell = sheet.cell(HEADER_ROW, column)
        _set_value(cell, header)
        cell.font = Font(name="Yu Gothic", bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor=_NAVY)
        cell.alignment = Alignment(vertical="center", wrap_text=True)
    sheet.row_dimensions[HEADER_ROW].height = 28

    border = Border(bottom=Side(style="hair", color="D5DBE5"))
    for index, values in enumerate(_report_rows(sheet_changes, cell_changes), start=1):
        row_number = HEADER_ROW + index
        fill = PatternFill("solid", fgColor=_ROW_COLORS[values[0]])
        for column, value in enumerate((index, *values), start=1):
            cell = sheet.cell(row_number, column)
            _set_value(cell, value)
            cell.font = Font(name="Yu Gothic", size=11, color="243746")
            cell.fill = fill
            cell.border = border
            cell.alignment = Alignment(vertical="top", wrap_text=True)

        # 長い値は折り返し表示。巨大な値でも行高が画面を占有しすぎないようにする。
        text_lines = max(
            (sum(max(1, (len(line) + 23) // 24) for line in str(value or "").split("\n"))
             for value in values),
            default=1,
        )
        sheet.row_dimensions[row_number].height = min(150, max(30, 16 * text_lines + 10))

    for column, width in zip("ABCDEFG", (12, 16, 24, 18, 38, 38, 34)):
        sheet.column_dimensions[column].width = width
    last_row = HEADER_ROW + count
    sheet.auto_filter.ref = f"A{HEADER_ROW}:G{last_row}"
    sheet.print_title_rows = f"1:{HEADER_ROW}"
    sheet.print_area = f"A1:G{last_row}"
    sheet.page_setup.orientation = "landscape"
    sheet.page_setup.paperSize = sheet.PAPERSIZE_A4
    sheet.page_setup.fitToWidth = 1
    sheet.page_setup.fitToHeight = 0
    return workbook


def write_report(sheet_changes, cell_changes, old_file: Path, new_file: Path,
                 output_path: Path | None = None) -> Path:
    """差分を新しい.xlsxへ保存し、絶対パスを返す。既存ファイルは上書きしない。"""
    created_at = datetime.now()
    output = Path(output_path) if output_path is not None else (
        REPORTS_DIR / f"変更点まとめ_{created_at:%Y%m%d_%H%M%S_%f}.xlsx"
    )
    workbook = None
    created_file = False
    try:
        output = output.expanduser().resolve()
        old_file, new_file = Path(old_file), Path(new_file)
        if output.suffix.lower() != ".xlsx":
            raise ReportWriteError("出力ファイルの拡張子は .xlsx を指定してください。")
        if output in (old_file.resolve(), new_file.resolve()):
            raise ReportWriteError("入力ファイルと同じ場所には出力できません。")
        if output.exists():
            raise ReportWriteError(f"出力先がすでに存在します: {output}")

        workbook = _build_workbook(
            sheet_changes, cell_changes, old_file, new_file, created_at
        )
        output.parent.mkdir(parents=True, exist_ok=True)
        # 排他作成により、確認後に同じ名前が作成された場合も上書きしない。
        with output.open("xb") as stream:
            created_file = True
            workbook.save(stream)
    except FileExistsError as exc:
        raise ReportWriteError(f"出力先がすでに存在します: {output}") from exc
    except (OSError, ValueError, TypeError, IllegalCharacterError, OverflowError) as exc:
        if created_file:
            try:
                output.unlink()
            except OSError:
                pass
        raise ReportWriteError(f"Excelレポートを保存できません: {output} ({exc})") from exc
    finally:
        if workbook is not None:
            workbook.close()
    return output


# ============================================================
# メイン処理
# ============================================================

def configure_output_encoding():
    """日本語の結果を、OSの既定文字コードに左右されず出力する。"""
    for stream in (sys.stdout, sys.stderr):
        if hasattr(stream, "reconfigure"):
            stream.reconfigure(encoding="utf-8")


def main(argv=None):
    configure_output_encoding()
    parser = argparse.ArgumentParser(
        description="Excelファイルのシート構成とセルの値・数式を比較します。",
        epilog="引数を省略すると、同梱のexamples/old.xlsxとexamples/new.xlsxを比較します。",
    )
    parser.add_argument("old_file", nargs="?", type=Path, help="変更前のExcelファイル")
    parser.add_argument("new_file", nargs="?", type=Path, help="変更後のExcelファイル")
    parser.add_argument(
        "-o", "--output", type=Path, metavar="PATH",
        help="Excelレポートの保存先（省略時: reports/変更点まとめ_日時.xlsx）",
    )
    args = parser.parse_args(argv)

    if (args.old_file is None) != (args.new_file is None):
        parser.error("変更前・変更後のファイルを両方指定してください。")

    old_file = args.old_file or EXAMPLES_DIR / "old.xlsx"
    new_file = args.new_file or EXAMPLES_DIR / "new.xlsx"

    # 後のファイルの読み込みや比較に失敗した場合も、読み込んだブックを閉じる。
    with ExitStack() as stack:
        try:
            old_wb = stack.enter_context(closing(load_workbook(old_file)))
            new_wb = stack.enter_context(closing(load_workbook(new_file)))
        except WorkbookReadError as exc:
            print(f"エラー: {exc}", file=sys.stderr)
            return 1

        sheet_changes = detect_sheet_changes(old_wb, new_wb)
        cell_changes = []
        new_sheet_set = set(new_wb.sheetnames)

        # 両方に存在するシートだけ比較する。
        for sheet_name in old_wb.sheetnames:
            if sheet_name not in new_sheet_set:
                continue

            cell_changes.extend(compare_sheet(
                sheet_name,
                old_wb[sheet_name],
                new_wb[sheet_name],
            ))

        print_results(sheet_changes, cell_changes)

    try:
        report_path = write_report(
            sheet_changes, cell_changes, old_file, new_file,
            output_path=args.output,
        )
    except ReportWriteError as exc:
        print(f"エラー: {exc}", file=sys.stderr)
        return 1

    print(f"レポート保存先: {report_path}")
    return 0


if __name__ == "__main__":
    sys.exit(main())
